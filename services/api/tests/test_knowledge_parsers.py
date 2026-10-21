from io import BytesIO
from pathlib import Path

from docx import Document
from pptx import Presentation
from reportlab.pdfgen import canvas

from app.knowledge import KnowledgeStore
from app.storage import ProjectStore


def _project(root: Path):
    return {
        "id": "knowledge-project", "name": "Knowledge QA", "slug": "knowledge-qa", "brief": "local knowledge", "budgetMinor": None, "currency": "CNY", "status": "draft", "rootPath": str(root / "knowledge-qa"), "createdAt": "2026-09-01T00:00:00+00:00", "updatedAt": "2026-09-01T00:00:00+00:00",
    }


def test_pdf_docx_pptx_are_ingested(tmp_path):
    store = ProjectStore(str(tmp_path))
    store.create_project(_project(tmp_path))
    knowledge = KnowledgeStore(store)

    pdf_buffer = BytesIO()
    pdf = canvas.Canvas(pdf_buffer)
    pdf.drawString(72, 760, "PDF research source")
    pdf.save()
    pdf_result = knowledge.ingest_file(project_id="knowledge-project", filename="research.pdf", content=pdf_buffer.getvalue())
    assert pdf_result["parser"] == "pypdf+pdfium"
    assert pdf_result["chunkCount"] == 1
    assert pdf_result["pageCount"] == 1
    assert pdf_result["partCount"] == 1

    docx_buffer = BytesIO()
    document = Document()
    document.add_paragraph("DOCX craft notes")
    document.save(docx_buffer)
    docx_result = knowledge.ingest_file(project_id="knowledge-project", filename="craft.docx", content=docx_buffer.getvalue())
    assert docx_result["parser"] == "python-docx"
    assert docx_result["chunkCount"] == 1

    pptx_buffer = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX competitor notes"
    presentation.save(pptx_buffer)
    pptx_result = knowledge.ingest_file(project_id="knowledge-project", filename="competitor.pptx", content=pptx_buffer.getvalue())
    assert pptx_result["parser"] == "python-pptx"
    assert pptx_result["chunkCount"] == 1

    documents = knowledge.list_documents("knowledge-project")
    assert len(documents) == 3
    assert all((tmp_path / "knowledge-qa" / item["pathRelative"]).is_file() for item in documents)


def test_image_ingestion_keeps_local_ocr_fallback(monkeypatch):
    monkeypatch.setattr("app.knowledge.shutil.which", lambda _name: None)
    text, parser, warning = KnowledgeStore._extract_text(".png", b"not-a-real-image")
    assert text == ""
    assert parser == "image-metadata-only"
    assert "Tesseract" in warning
