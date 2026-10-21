from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable


PARSER_VERSION = 2
MAX_PAGES = 1000
MAX_IMAGES = 1000
MAX_OOXML_MEMBERS = 2000
MAX_OOXML_UNCOMPRESSED = 512 * 1024 * 1024
MAX_PART_BYTES = 250 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
PREVIEW_MAX_EDGE = 2048
PREVIEW_MAX_BYTES = 8 * 1024 * 1024


@dataclass
class ParsedDocument:
    text: str
    parser: str
    blocks: list[dict[str, Any]] = field(default_factory=list)
    parts: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    page_count: int = 0

    @property
    def image_count(self) -> int:
        return sum(1 for part in self.parts if part.get("kind") == "embedded_image")

    @property
    def status(self) -> str:
        if self.warnings and (self.text.strip() or self.parts):
            return "partial"
        if self.text.strip() or self.parts:
            return "ready"
        return "stored"

    @property
    def warning(self) -> str | None:
        messages = list(dict.fromkeys(message.strip() for message in self.warnings if message.strip()))
        return "；".join(messages) if messages else None


def _clean_text(value: Any) -> str:
    return re.sub(r"[ \t]+", " ", str(value or "")).strip()


def _safe_name(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9._-]+", "-", Path(value or fallback).name).strip("-.")
    return cleaned or fallback


def _preflight_ooxml(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_OOXML_MEMBERS:
                raise ValueError(f"Office 文档包含超过 {MAX_OOXML_MEMBERS} 个压缩成员，已拒绝解析。")
            if sum(max(0, member.file_size) for member in members) > MAX_OOXML_UNCOMPRESSED:
                raise ValueError("Office 文档解压后超过 512 MB，已拒绝解析。")
            if any(member.flag_bits & 0x1 for member in members):
                raise ValueError("加密 Office 文档无法在本地解析。")
    except zipfile.BadZipFile as error:
        raise ValueError("Office 文档结构损坏，已保留原文件。") from error


def _normalized_image(content: bytes, filename: str, declared_mime: str | None = None) -> dict[str, Any]:
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        with Image.open(BytesIO(content)) as opened:
            width, height = opened.size
            if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                raise ValueError("图片像素尺寸超过安全限制")
            image = ImageOps.exif_transpose(opened)
            image.seek(0)
            image.load()
            has_alpha = image.mode in {"RGBA", "LA"} or "transparency" in image.info
            if has_alpha:
                canonical = image.convert("RGBA")
                extension, mime, image_format = ".png", "image/png", "PNG"
            else:
                canonical = image.convert("RGB")
                if str(declared_mime or "").lower() == "image/jpeg" or Path(filename).suffix.lower() in {".jpg", ".jpeg"}:
                    extension, mime, image_format = ".jpg", "image/jpeg", "JPEG"
                else:
                    extension, mime, image_format = ".png", "image/png", "PNG"
            source_buffer = BytesIO()
            canonical.save(source_buffer, image_format, quality=95, optimize=True)
            preview = canonical.copy()
            preview.thumbnail((PREVIEW_MAX_EDGE, PREVIEW_MAX_EDGE))
            preview_buffer = BytesIO()
            preview_format = "PNG" if has_alpha else "JPEG"
            preview.save(preview_buffer, preview_format, quality=88, optimize=True)
            preview_bytes = preview_buffer.getvalue()
            if len(preview_bytes) > PREVIEW_MAX_BYTES:
                preview.thumbnail((1280, 1280))
                preview_buffer = BytesIO()
                preview.convert("RGB").save(preview_buffer, "JPEG", quality=82, optimize=True)
                preview_bytes = preview_buffer.getvalue()
            return {
                "content": source_buffer.getvalue(),
                "preview": preview_bytes,
                "mime": mime,
                "previewMime": "image/png" if preview_format == "PNG" else "image/jpeg",
                "filename": _safe_name(Path(filename).stem, "image") + extension,
                "width": canonical.width,
                "height": canonical.height,
                "availability": "available",
            }
    except (UnidentifiedImageError, OSError, ValueError) as error:
        raise ValueError(str(error) or "图片格式无法读取") from error


def _append_image(
    parsed: ParsedDocument,
    *,
    content: bytes,
    filename: str,
    mime: str | None,
    locator: str,
    block_index: int | None,
    page_number: int | None = None,
    slide_number: int | None = None,
    metadata: dict[str, Any] | None = None,
) -> None:
    if parsed.image_count >= MAX_IMAGES:
        if not any("图片数量超过" in warning for warning in parsed.warnings):
            parsed.warnings.append(f"图片数量超过 {MAX_IMAGES} 张，后续图片未提取")
        return
    try:
        image = _normalized_image(content, filename, mime)
    except ValueError as error:
        parsed.parts.append({
            "kind": "embedded_image",
            "locator": locator,
            "blockIndex": block_index,
            "pageNumber": page_number,
            "slideNumber": slide_number,
            "filename": _safe_name(filename, "unsupported-image"),
            "mime": mime or "application/octet-stream",
            "availability": "unsupported",
            "metadata": {**(metadata or {}), "reason": str(error)},
        })
        parsed.warnings.append(f"{locator} 的图片无法读取：{error}")
        return
    parsed.parts.append({
        "kind": "embedded_image",
        "locator": locator,
        "blockIndex": block_index,
        "pageNumber": page_number,
        "slideNumber": slide_number,
        "metadata": metadata or {},
        **image,
    })


def _docx_images(element: Any, related_parts: dict[str, Any]) -> Iterable[tuple[str, Any]]:
    from docx.oxml.ns import qn

    for index, blip in enumerate(element.xpath(".//a:blip"), start=1):
        relation_id = blip.get(qn("r:embed"))
        part = related_parts.get(relation_id)
        if relation_id and part is not None and hasattr(part, "blob"):
            yield f"{relation_id}:{index}", part


def _parse_docx(content: bytes, *, include_parts: bool) -> ParsedDocument:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _preflight_ooxml(content)
    document = Document(BytesIO(content))
    parsed = ParsedDocument(text="", parser="python-docx")

    def add_container(container: Any, element: Any, kind: str, locator: str) -> None:
        text = _clean_text(container.text if hasattr(container, "text") else "")
        images = list(_docx_images(element, container.part.related_parts)) if include_parts else []
        block_index = len(parsed.blocks) if text or images else None
        if text or images:
            parsed.blocks.append({"kind": kind, "text": text, "locator": locator, "metadata": {}})
        if not include_parts:
            return
        for image_index, (_relation_key, image_part) in enumerate(images, start=1):
            _append_image(
                parsed,
                content=image_part.blob,
                filename=Path(str(image_part.partname)).name,
                mime=getattr(image_part, "content_type", None),
                locator=f"{locator}/image:{image_index}",
                block_index=block_index,
                metadata={"officePart": str(image_part.partname)},
            )

    body_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            body_index += 1
            paragraph = Paragraph(child, document)
            add_container(paragraph, paragraph._p, "paragraph", f"body/paragraph:{body_index}")
        elif child.tag == qn("w:tbl"):
            body_index += 1
            table = Table(child, document)
            rows = [" | ".join(_clean_text(cell.text) for cell in row.cells) for row in table.rows]
            text = "\n".join(row for row in rows if row.strip(" |"))
            images = list(_docx_images(table._tbl, table.part.related_parts)) if include_parts else []
            block_index = len(parsed.blocks) if text or images else None
            if text or images:
                parsed.blocks.append({"kind": "table", "text": text, "locator": f"body/table:{body_index}", "metadata": {"rows": len(table.rows)}})
            if include_parts:
                for image_index, (_key, image_part) in enumerate(images, start=1):
                    _append_image(parsed, content=image_part.blob, filename=Path(str(image_part.partname)).name,
                                  mime=getattr(image_part, "content_type", None), locator=f"body/table:{body_index}/image:{image_index}",
                                  block_index=block_index, metadata={"officePart": str(image_part.partname)})

    seen_story_parts: set[str] = set()
    for section_number, section in enumerate(document.sections, start=1):
        for story_name, story in (("header", section.header), ("footer", section.footer)):
            part_name = str(story.part.partname)
            if part_name in seen_story_parts:
                continue
            seen_story_parts.add(part_name)
            for paragraph_number, paragraph in enumerate(story.paragraphs, start=1):
                add_container(paragraph, paragraph._p, story_name, f"section:{section_number}/{story_name}/paragraph:{paragraph_number}")
            for table_number, table in enumerate(story.tables, start=1):
                rows = [" | ".join(_clean_text(cell.text) for cell in row.cells) for row in table.rows]
                text = "\n".join(row for row in rows if row.strip(" |"))
                images = list(_docx_images(table._tbl, table.part.related_parts)) if include_parts else []
                block_index = len(parsed.blocks) if text or images else None
                if text or images:
                    parsed.blocks.append({"kind": f"{story_name}_table", "text": text,
                                          "locator": f"section:{section_number}/{story_name}/table:{table_number}", "metadata": {}})
                for image_index, (_key, image_part) in enumerate(images, start=1):
                    _append_image(parsed, content=image_part.blob, filename=Path(str(image_part.partname)).name,
                                  mime=getattr(image_part, "content_type", None),
                                  locator=f"section:{section_number}/{story_name}/table:{table_number}/image:{image_index}",
                                  block_index=block_index, metadata={"officePart": str(image_part.partname)})

    parsed.text = "\n".join(block["text"] for block in parsed.blocks if block.get("text"))
    return parsed


def _parse_pptx(content: bytes, *, include_parts: bool) -> ParsedDocument:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    _preflight_ooxml(content)
    presentation = Presentation(BytesIO(content))
    parsed = ParsedDocument(text="", parser="python-pptx", page_count=min(len(presentation.slides), MAX_PAGES))
    if len(presentation.slides) > MAX_PAGES:
        parsed.warnings.append(f"幻灯片超过 {MAX_PAGES} 页，后续页面未解析")

    def walk_shapes(shapes: Any, slide_number: int, prefix: str, slide_block: int) -> None:
        for order, shape in enumerate(shapes, start=1):
            shape_id = getattr(shape, "shape_id", order)
            locator = f"slide:{slide_number}/{prefix}shape:{shape_id}"
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk_shapes(shape.shapes, slide_number, f"{prefix}group:{shape_id}/", slide_block)
                continue
            if getattr(shape, "has_table", False):
                rows = [" | ".join(_clean_text(cell.text) for cell in row.cells) for row in shape.table.rows]
                text = "\n".join(row for row in rows if row.strip(" |"))
                if text:
                    parsed.blocks.append({"kind": "table", "text": text, "locator": locator,
                                          "metadata": {"slideNumber": slide_number, "shapeName": getattr(shape, "name", "")}})
            elif getattr(shape, "has_text_frame", False):
                text = _clean_text(shape.text)
                if text:
                    parsed.blocks.append({"kind": "slide_text", "text": text, "locator": locator,
                                          "metadata": {"slideNumber": slide_number, "shapeName": getattr(shape, "name", "")}})
            if include_parts and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                _append_image(parsed, content=image.blob, filename=getattr(image, "filename", f"slide-{slide_number}.png"),
                              mime=getattr(image, "content_type", None), locator=f"{locator}/image",
                              block_index=slide_block, slide_number=slide_number,
                              metadata={"shapeName": getattr(shape, "name", ""), "left": int(shape.left), "top": int(shape.top),
                                        "width": int(shape.width), "height": int(shape.height),
                                        "crop": {key: float(getattr(shape, key, 0) or 0) for key in ("crop_left", "crop_top", "crop_right", "crop_bottom")}})

    for slide_number, slide in enumerate(presentation.slides, start=1):
        if slide_number > MAX_PAGES:
            break
        slide_block = len(parsed.blocks)
        parsed.blocks.append({"kind": "slide", "text": f"幻灯片 {slide_number}", "locator": f"slide:{slide_number}",
                              "metadata": {"slideNumber": slide_number}})
        walk_shapes(slide.shapes, slide_number, "", slide_block)
        try:
            notes = _clean_text(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide else ""
        except (AttributeError, ValueError):
            notes = ""
        if notes:
            parsed.blocks.append({"kind": "notes", "text": notes, "locator": f"slide:{slide_number}/notes",
                                  "metadata": {"slideNumber": slide_number}})
    parsed.text = "\n".join(block["text"] for block in parsed.blocks if block.get("text"))
    return parsed


def _parse_pdf(content: bytes, *, include_parts: bool) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted and not reader.decrypt(""):
        return ParsedDocument(text="", parser="pypdf+pdfium", warnings=["加密 PDF 无法解析，已保留原文件"])
    parsed = ParsedDocument(text="", parser="pypdf+pdfium", page_count=min(len(reader.pages), MAX_PAGES))
    if len(reader.pages) > MAX_PAGES:
        parsed.warnings.append(f"PDF 超过 {MAX_PAGES} 页，后续页面未解析")
    for page_number, page in enumerate(reader.pages[:MAX_PAGES], start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as error:
            text = ""
            parsed.warnings.append(f"第 {page_number} 页文本提取失败：{error}")
        block_index = len(parsed.blocks)
        parsed.blocks.append({"kind": "page", "text": text, "locator": f"page:{page_number}",
                              "metadata": {"pageNumber": page_number}})
        if include_parts:
            parsed.parts.append({"kind": "page_preview", "locator": f"page:{page_number}/preview", "blockIndex": block_index,
                                 "pageNumber": page_number, "filename": f"page-{page_number}.png", "mime": "image/png",
                                 "availability": "virtual", "metadata": {"renderOnDemand": True}})
            try:
                images = list(page.images)
            except Exception as error:
                images = []
                parsed.warnings.append(f"第 {page_number} 页内嵌图片读取失败：{error}")
            for image_number, image in enumerate(images, start=1):
                _append_image(parsed, content=image.data, filename=getattr(image, "name", f"page-{page_number}-{image_number}.png"),
                              mime=None, locator=f"page:{page_number}/image:{image_number}", block_index=block_index,
                              page_number=page_number, metadata={})
    parsed.text = "\n".join(block["text"] for block in parsed.blocks if block.get("text"))
    if not parsed.text.strip():
        parsed.warnings.append("PDF 没有可提取文本；可通过页面预览进行视觉读取")
    return parsed


def parse_document(suffix: str, content: bytes, *, include_parts: bool = True) -> ParsedDocument:
    suffix = suffix.lower()
    if suffix == ".pdf":
        try:
            return _parse_pdf(content, include_parts=include_parts)
        except ImportError:
            return ParsedDocument(text="", parser="unavailable", warnings=["未安装 PDF 解析依赖，已保存原文件"])
        except Exception as error:
            return ParsedDocument(text="", parser="pypdf+pdfium", warnings=[f"PDF 解析失败：{error}"])
    if suffix == ".docx":
        try:
            return _parse_docx(content, include_parts=include_parts)
        except ImportError:
            return ParsedDocument(text="", parser="unavailable", warnings=["未安装 python-docx，已保存原文件"])
        except Exception as error:
            return ParsedDocument(text="", parser="python-docx", warnings=[str(error)])
    if suffix == ".pptx":
        try:
            return _parse_pptx(content, include_parts=include_parts)
        except ImportError:
            return ParsedDocument(text="", parser="unavailable", warnings=["未安装 python-pptx，已保存原文件"])
        except Exception as error:
            return ParsedDocument(text="", parser="python-pptx", warnings=[str(error)])
    if suffix in {".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm"}:
        text = content.decode("utf-8", errors="replace")
        return ParsedDocument(text=text, parser="text", blocks=[{"kind": "text", "text": text, "locator": "body", "metadata": {}}])
    return ParsedDocument(text="", parser="binary-metadata-only", warnings=["该文件类型暂不支持结构化提取，已保存原文件"])


def render_pdf_page(path: Path, page_number: int) -> tuple[bytes, int, int]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        if page_number < 1 or page_number > len(document):
            raise ValueError("PDF 页码不存在")
        page = document[page_number - 1]
        try:
            width, height = page.get_size()
            scale = min(2.0, PREVIEW_MAX_EDGE / max(float(width), float(height), 1.0))
            bitmap = page.render(scale=max(scale, 0.5), rev_byteorder=True)
            image = bitmap.to_pil()
            image.thumbnail((PREVIEW_MAX_EDGE, PREVIEW_MAX_EDGE))
            output = BytesIO()
            image.save(output, "PNG", optimize=True)
            return output.getvalue(), image.width, image.height
        finally:
            page.close()
    finally:
        document.close()


def stable_part_id(document_id: str, locator: str, checksum: str = "") -> str:
    return hashlib.sha256(f"{document_id}\0{locator}\0{checksum}".encode()).hexdigest()[:32]
