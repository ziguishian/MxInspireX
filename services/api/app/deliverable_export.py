"""Export current canvas deliverables without requiring the retired report flow."""
from __future__ import annotations

import base64
import html
import io
import json
import zipfile
from pathlib import Path
from uuid import uuid4
from pydantic import BaseModel, Field
from .html_assets import resolve_html_assets


class Slide(BaseModel):
    title: str = Field(max_length=200)
    body: str = Field(default='', max_length=16000)
    imageAssetIds: list[str] = Field(default_factory=list, max_length=8)


def export_deliverable(registry, *, project_id, task_id, format, title='设计交付', node_ids=None, slides=None, **kwargs):
    if format not in {'html','pdf','pptx','zip'}: raise ValueError('支持 HTML、PDF、PPTX、ZIP')
    board=registry.canvas_store.load(project_id)
    ids=set(node_ids or [])
    nodes=[] if slides and not ids else [n for n in board['nodes'] if (n['id'] in ids if ids else not n['data'].get('downloadable'))]
    if ids - {n['id'] for n in nodes}: raise ValueError('导出节点不存在或不属于项目')
    if not nodes and not slides: raise ValueError('没有可导出的内容')
    if any(n['data'].get('generating') or n['data'].get('error') for n in nodes): raise ValueError('请等待生成完成并修复失败节点后导出，或明确选择已完成节点')
    def node_body(node):
        data=node['data']
        if data.get('kind')=='html' and data.get('html'):
            from .research import ResearchStore
            return ResearchStore._extract(data['html'].encode(),'text/html')[0]
        return data.get('text') or data.get('summary') or (f"音视频素材：{data.get('title','素材')}。完整媒体请使用 HTML 或 ZIP 导出。" if data.get('kind') in {'audio','video'} else '')
    authored=[Slide.model_validate(s) for s in slides] if slides else [Slide(title=n['data'].get('title','设计内容'),body=node_body(n),imageAssetIds=[n['data']['assetId']] if n['data'].get('kind')=='image' and n['data'].get('assetId') else []) for n in nodes]
    if not 1 <= len(authored) <= 80: raise ValueError('导出需要 1–80 页内容')
    assets={}
    for asset_id in {a for s in authored for a in s.imageAssetIds} | {n['data']['assetId'] for n in nodes if n['data'].get('assetId')}:
        asset=registry.asset_store.get_asset(project_id=project_id,asset_id=asset_id)
        if not asset or not Path(asset['absolutePath']).is_file(): raise ValueError('导出引用的资产不存在')
        assets[asset_id]=asset
    if task_id:
        from .svg_policy import contains_svg, ERROR
        if any(asset.get('mime') == 'image/svg+xml' or Path(asset['absolutePath']).suffix.lower() == '.svg' for asset in assets.values()) or any(contains_svg(n['data'].get('html') or n['data'].get('text', '')) for n in nodes):
            raise ValueError(ERROR + '请选择实际生成的栅格图像与排版文档导出。')
    from .cad import cad_manifest
    for asset in list(assets.values()):
        if asset.get('metadata',{}).get('cadRole') == 'model':
            model = cad_manifest(registry.asset_store, project_id, asset['id'])
            for child_id in model['files'].values():
                assets[child_id] = registry.asset_store.get_asset(project_id=project_id,asset_id=child_id)
    for slide in authored:
        if any(not assets[a]['mime'].startswith('image/') for a in slide.imageAssetIds): raise ValueError('演示页图像引用必须为图片')
    document='<meta charset="utf-8"><title>'+html.escape(title)+'</title><style>body{font:18px sans-serif;max-width:1000px;margin:40px auto;line-height:1.6}section{break-after:page;margin:3rem 0}img{max-width:100%;max-height:600px}pre{white-space:pre-wrap}</style>'
    for slide in authored:
        document+='<section><h1>'+html.escape(slide.title)+'</h1><pre>'+html.escape(slide.body)+'</pre>'
        for asset_id in slide.imageAssetIds:
            asset=assets[asset_id]
            document+=f'<img alt="{html.escape(slide.title,quote=True)}" src="data:{asset["mime"]};base64,{base64.b64encode(Path(asset["absolutePath"]).read_bytes()).decode()}">'
        document+='</section>'
    for node in nodes:
        data=node['data'];kind=data.get('kind')
        if kind in {'audio','video'} and data.get('assetId'):
            asset=assets[data['assetId']]
            document+=f'<section><h2>{html.escape(data.get("title","媒体"))}</h2><{kind} controls src="data:{asset["mime"]};base64,{base64.b64encode(Path(asset["absolutePath"]).read_bytes()).decode()}"></{kind}></section>'
    stream=io.BytesIO()
    if format=='html':
        # A single HTML node exports its actual source, preserving interactions.
        source=resolve_html_assets(nodes[0]['data']) if len(nodes)==1 and nodes[0]['data'].get('kind')=='html' and not slides else document
        stream.write(source.encode())
    elif format=='zip':
        with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('index.html',document)
            archive.writestr('manifest.json',json.dumps({'title':title,'nodes':nodes,'slides':[s.model_dump() for s in authored]},ensure_ascii=False,indent=2))
            for index,node in enumerate(nodes):
                data=node['data']
                if data.get('kind') in {'html','text'}: archive.writestr(f'nodes/{index+1}.{ "html" if data["kind"]=="html" else "txt"}',resolve_html_assets(data) if data['kind']=='html' else data.get('text') or data.get('summary',''))
            for asset in assets.values(): archive.write(asset['absolutePath'],f'assets/{asset["id"]}{Path(asset["absolutePath"]).suffix}')
    elif format=='pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from PIL import Image
        deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        for spec in authored:
            # Split long text rather than silently clipping it from delivery.
            chunk_size=320 if spec.imageAssetIds else 700
            parts=[spec.body[i:i+chunk_size] for i in range(0,len(spec.body),chunk_size)] or ['']
            for part_index,body in enumerate(parts):
                slide=deck.slides.add_slide(deck.slide_layouts[6])
                heading=slide.shapes.add_textbox(Inches(.6),Inches(.4),Inches(12),Inches(.7)).text_frame
                heading.text=spec.title + ('（续）' if part_index else '');heading.paragraphs[0].font.size=Pt(28)
                pictures=spec.imageAssetIds if part_index==0 else []
                box=slide.shapes.add_textbox(Inches(.6),Inches(1.4),Inches(6 if pictures else 12),Inches(5.5)).text_frame
                box.word_wrap=True;box.text=body
                for paragraph in box.paragraphs: paragraph.font.size=Pt(18);paragraph.font.name='Microsoft YaHei'
                for i,asset_id in enumerate(pictures):
                    path=assets[asset_id]['absolutePath']
                    with Image.open(path) as image:
                        image=image.convert('RGB');out=io.BytesIO();image.save(out,format='PNG');out.seek(0)
                        width,height=image.size
                    cell_h=5.2/len(pictures);scale=min(5.5/width,cell_h/height)
                    slide.shapes.add_picture(out,Inches(7),Inches(1.4+i*cell_h),width=Inches(width*scale),height=Inches(height*scale))
        deck.save(stream)
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        styles=getSampleStyleSheet()
        for name in ['Normal','Title']: styles[name].fontName='STSong-Light'
        styles['Normal'].fontSize=11;styles['Normal'].leading=17;styles['Normal'].wordWrap='CJK'
        flow=[]
        for spec in authored:
            if flow: flow.append(PageBreak())
            flow += [Paragraph(html.escape(spec.title),styles['Title']),Spacer(1,15)]
            flow += [Paragraph(html.escape(line) or ' ',styles['Normal']) for line in spec.body.splitlines()]
            for asset_id in spec.imageAssetIds:
                picture=Image(assets[asset_id]['absolutePath']);picture._restrictSize(480,480);flow += [Spacer(1,12),picture]
        SimpleDocTemplate(stream).build(flow)
    asset=registry.asset_store.add_bytes(project_id=project_id,kind='document',filename=f'deliverable-{uuid4().hex}.{format}',content=stream.getvalue(),metadata={'taskId':task_id,'source':'agent-export','executionReceipt':kwargs.get('_execution_state',{}).get('fingerprint'),'title':title,'sourceNodeIds':[n['id'] for n in nodes]})
    node=registry.canvas_store.create_node(project_id,kind='text',title=title,task_id=task_id,execution_receipt=kwargs.get('_execution_state',{}).get('fingerprint'),data={'text':f'{title} · {format.upper()} · {len(authored)} 页/项','assetId':asset['id'],'asset':{k:v for k,v in asset.items() if k!='absolutePath'},'downloadable':True})
    return {'status':'succeeded','assetId':asset['id'],'asset':asset,'node':node,'format':format}
